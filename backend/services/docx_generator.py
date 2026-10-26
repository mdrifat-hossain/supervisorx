import io
import re
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from services.document_models import StructuredDocument, ContentBlock, SectionModel


# Professional Academic Color Palette
COLOR_NAVY = RGBColor(15, 41, 66)       # #0F2942
COLOR_TEAL = RGBColor(2, 132, 199)      # #0284C7
COLOR_TEXT = RGBColor(30, 41, 59)       # #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B
HEX_HEADER_BG = "0F2942"
HEX_ROW_ALT = "F8FAFC"
HEX_BORDER = "CBD5E1"


def set_cell_background(cell, fill_hex: str):
    """Sets background fill color for a Word table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Sets cell padding (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def _extract_braced_group(s: str, start: int):
    """Extracts a {...} group from string s starting at index start. Returns (content, end_index)"""
    if start >= len(s) or s[start] != '{':
        return None, start
    depth = 0
    content = []
    for idx in range(start, len(s)):
        ch = s[idx]
        if ch == '{':
            depth += 1
            if depth > 1:
                content.append(ch)
        elif ch == '}':
            depth -= 1
            if depth == 0:
                return "".join(content), idx + 1
            else:
                content.append(ch)
        else:
            content.append(ch)
    return None, start


def clean_latex_for_doc(latex: str) -> str:
    """Converts LaTeX notation into clean, readable scientific typography for document text."""
    clean = latex.strip()

    # Safely replace \frac{num}{den} with matching braces up to 10 passes
    for _ in range(10):
        pos = clean.find(r"\frac")
        if pos == -1:
            break
        num_start = pos + len(r"\frac")
        while num_start < len(clean) and clean[num_start].isspace():
            num_start += 1
        num, den_start = _extract_braced_group(clean, num_start)
        if num is None:
            clean = clean[:pos] + clean[pos + len(r"\frac"):]
            continue
        while den_start < len(clean) and clean[den_start].isspace():
            den_start += 1
        den, end_pos = _extract_braced_group(clean, den_start)
        if den is None:
            clean = clean[:pos] + f"({num})" + clean[den_start:]
            continue
        clean = clean[:pos] + f"({num} / {den})" + clean[end_pos:]

    # Replace common symbols
    replacements = [
        (r"\\left\(", "("),
        (r"\\right\)", ")"),
        (r"\\left\[", "["),
        (r"\\right\]", "]"),
        (r"\\left\{", "{"),
        (r"\\right\}", "}"),
        (r"\\cdot", " · "),
        (r"\\times", " × "),
        (r"\\pm", " ± "),
        (r"\\leq", " ≤ "),
        (r"\\geq", " ≥ "),
        (r"\\neq", " ≠ "),
        (r"\\approx", " ≈ "),
        (r"\\infty", "∞"),
        (r"\\alpha", "α"),
        (r"\\beta", "β"),
        (r"\\gamma", "γ"),
        (r"\\delta", "δ"),
        (r"\\Delta", "Δ"),
        (r"\\epsilon", "ε"),
        (r"\\lambda", "λ"),
        (r"\\mu", "μ"),
        (r"\\sigma", "σ"),
        (r"\\theta", "θ"),
        (r"\\omega", "ω"),
        (r"\\Omega", "Ω"),
        (r"\\sum", "∑"),
        (r"\\prod", "∏"),
        (r"\\int", "∫"),
        (r"\\partial", "∂"),
        (r"\\sqrt\{([^}]+)\}", r"√(\1)"),
    ]
    for pattern, rep in replacements:
        clean = re.sub(pattern, rep, clean)

    # Clean redundant braces
    clean = re.sub(r"\{([^{}]+)\}", r"\1", clean)
    clean = clean.replace("\\", "")
    return clean


def add_formatted_runs(paragraph, text: str):
    """Parses inline bold (**text**), code (`text`), inline math ($math$), and citations ([1])."""
    tokens = re.split(r"(\[\d+(?:,\s*\d+)*\]|\$[^$\n]+\$|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for token in tokens:
        if not token:
            continue
        if re.match(r"^\[\d+(?:,\s*\d+)*\]$", token):
            # Citation badge
            run = paragraph.add_run(f" {token}")
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_TEAL
        elif token.startswith("$") and token.endsWith("$") if hasattr(token, "endsWith") else (token.startswith("$") and token.endswith("$") and len(token) > 2):
            # Inline math
            clean_math = clean_latex_for_doc(token[1:-1])
            run = paragraph.add_run(f" {clean_math} ")
            run.font.italic = True
            run.font.name = "Cambria Math"
            run.font.color.rgb = COLOR_NAVY
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_TEAL
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.font.italic = True
        else:
            paragraph.add_run(token)


def add_hyperlink(paragraph, text: str, url: str):
    """Adds a clickable external link to a DOCX paragraph."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0284C7")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def reference_pdf_url(ref) -> str:
    url = ref.url or ""
    if ref.page_start and "#page=" not in url:
        url = f"{url}#page={ref.page_start}"
    return url


def generate_docx_document(doc_model: StructuredDocument) -> bytes:
    """
    Renders a validated StructuredDocument into a polished, professional Microsoft Word (.docx) document.
    """
    doc = Document()

    # Set 1-inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)

        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"SupervisorX Scientific Report • {doc_model.project_name or 'Research Project'}")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run(f"Generated: {doc_model.generated_at} | Evidence-Grounded Research Synthesis")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_MUTED

    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run(doc_model.title)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_NAVY

    # Metadata & Scope Sub-bar
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(16)

    type_run = meta_p.add_run(f"{doc_model.document_type.replace('_', ' ').upper()}   |   ")
    type_run.font.size = Pt(9.5)
    type_run.font.bold = True
    type_run.font.color.rgb = COLOR_TEAL

    date_run = meta_p.add_run(f"Date: {doc_model.generated_at}")
    date_run.font.size = Pt(9.5)
    date_run.font.color.rgb = COLOR_MUTED

    if doc_model.scope_description:
        scope_p = doc.add_paragraph()
        scope_p.paragraph_format.space_before = Pt(0)
        scope_p.paragraph_format.space_after = Pt(14)
        scope_run = scope_p.add_run(f"Scope: {doc_model.scope_description}")
        scope_run.font.size = Pt(10)
        scope_run.font.italic = True
        scope_run.font.color.rgb = COLOR_MUTED

    # Horizontal Divider
    divider_p = doc.add_paragraph()
    divider_p.paragraph_format.space_before = Pt(0)
    divider_p.paragraph_format.space_after = Pt(14)
    d_run = divider_p.add_run("―" * 52)
    d_run.font.color.rgb = COLOR_MUTED
    d_run.font.size = Pt(9)

    # Render Sections
    for sec in doc_model.sections:
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(14 if sec.level == 1 else 10)
        h_p.paragraph_format.space_after = Pt(4)
        h_run = h_p.add_run(sec.heading)
        h_run.font.name = "Arial"
        h_run.font.bold = True

        if sec.level == 1:
            h_run.font.size = Pt(14)
            h_run.font.color.rgb = COLOR_NAVY
        elif sec.level == 2:
            h_run.font.size = Pt(12)
            h_run.font.color.rgb = COLOR_TEAL
        else:
            h_run.font.size = Pt(11)
            h_run.font.color.rgb = COLOR_TEXT

        # Blocks within Section
        for block in sec.blocks:
            if block.type == "paragraph" and block.content:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, block.content)

            elif block.type == "bullet_list" and block.items:
                for item in block.items:
                    lp = doc.add_paragraph(style='List Bullet')
                    lp.paragraph_format.space_before = Pt(1)
                    lp.paragraph_format.space_after = Pt(3)
                    lp.paragraph_format.line_spacing = 1.15
                    add_formatted_runs(lp, item)

            elif block.type == "numbered_list" and block.items:
                for idx, item in enumerate(block.items, start=1):
                    lp = doc.add_paragraph(style='List Number')
                    lp.paragraph_format.space_before = Pt(1)
                    lp.paragraph_format.space_after = Pt(3)
                    lp.paragraph_format.line_spacing = 1.15
                    add_formatted_runs(lp, item)

            elif block.type == "blockquote" and block.content:
                qp = doc.add_paragraph()
                qp.paragraph_format.left_indent = Inches(0.4)
                qp.paragraph_format.space_before = Pt(4)
                qp.paragraph_format.space_after = Pt(6)
                qrun = qp.add_run(f"“{block.content}”")
                qrun.font.italic = True
                qrun.font.color.rgb = COLOR_MUTED

            elif block.type == "equation" and block.latex:
                clean_formula = clean_latex_for_doc(block.latex)
                eq_table = doc.add_table(rows=1, cols=1)
                eq_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = eq_table.cell(0, 0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

                eq_p = cell.paragraphs[0]
                eq_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                eq_p.paragraph_format.space_before = Pt(4)
                eq_p.paragraph_format.space_after = Pt(4)
                er = eq_p.add_run(clean_formula)
                er.font.name = "Cambria Math"
                er.font.size = Pt(11)
                er.font.bold = True
                er.font.color.rgb = COLOR_NAVY

                doc.add_paragraph().paragraph_format.space_after = Pt(4)

            elif block.type == "table" and block.headers and block.rows:
                num_cols = len(block.headers)
                num_rows = len(block.rows) + 1
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Header Row
                hdr_cells = table.rows[0].cells
                for col_idx, header_text in enumerate(block.headers):
                    cell = hdr_cells[col_idx]
                    set_cell_background(cell, HEX_HEADER_BG)
                    set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                    hp = cell.paragraphs[0]
                    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    add_formatted_runs(hp, header_text)
                    for hrun in hp.runs:
                        hrun.font.bold = True
                        hrun.font.size = Pt(9.5)
                        hrun.font.color.rgb = RGBColor(255, 255, 255)

                # Body Rows
                for row_idx, row_data in enumerate(block.rows, start=1):
                    row_cells = table.rows[row_idx].cells
                    fill_hex = HEX_ROW_ALT if row_idx % 2 == 1 else "FFFFFF"
                    for col_idx in range(min(num_cols, len(row_data))):
                        cell = row_cells[col_idx]
                        set_cell_background(cell, fill_hex)
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        cp = cell.paragraphs[0]
                        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        add_formatted_runs(cp, row_data[col_idx])
                        for c_run in cp.runs:
                            c_run.font.size = Pt(9.5)
                            c_run.font.color.rgb = COLOR_TEXT

                doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # References Section
    if doc_model.references:
        ref_heading_p = doc.add_paragraph()
        ref_heading_p.paragraph_format.space_before = Pt(18)
        ref_heading_p.paragraph_format.space_after = Pt(6)
        rhrun = ref_heading_p.add_run("References & Grounded Citations")
        rhrun.font.name = "Arial"
        rhrun.font.size = Pt(14)
        rhrun.font.bold = True
        rhrun.font.color.rgb = COLOR_NAVY

        for ref in doc_model.references:
            rp = doc.add_paragraph()
            rp.paragraph_format.left_indent = Inches(0.3)
            rp.paragraph_format.space_before = Pt(2)
            rp.paragraph_format.space_after = Pt(4)

            num_run = rp.add_run(f"[{ref.index}] ")
            num_run.font.bold = True
            num_run.font.color.rgb = COLOR_TEAL

            title_run = rp.add_run(f"\"{ref.paper_title}\". ")
            title_run.font.bold = True
            title_run.font.color.rgb = COLOR_TEXT

            author_text = ref.authors.rstrip()
            if author_text and not author_text.endswith("."):
                author_text += "."
            author_run = rp.add_run(f"{author_text} ")
            author_run.font.color.rgb = COLOR_TEXT

            venue_str = f"{ref.venue}, " if ref.venue else ""
            year_str = f"{ref.year}. " if ref.year else ""
            location_parts = [part for part in (ref.section_title, ref.page_label) if part]
            page_str = f"Section: {ref.section_title}; {ref.page_label}" if len(location_parts) == 2 else (location_parts[0] if location_parts else "Location unavailable")
            details_run = rp.add_run(f"{venue_str}{year_str}{page_str}. ")
            details_run.font.italic = True
            details_run.font.color.rgb = COLOR_MUTED
            source_url = reference_pdf_url(ref)
            if source_url.startswith(("http://", "https://")):
                add_hyperlink(rp, "Open source PDF", source_url)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
